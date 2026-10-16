"""Genera emis-miles.xlsx: un estado financiero con el formato de una exportación tipo EMIS (cifras en miles de USD, un solo bloque
de hoja con Estado de Resultados, Balance y Flujo, periodos de más reciente a más antiguo). Empresa y cifras INVENTADAS."""
import datetime, json
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Company Financials'
Y = ['2024 ', '2023 ', '2022 ', '2021 ', '2020 ']

def row(label, *vals):
    ws.append([label, *vals])

ws.append(['', '', '', 'PROVEEDOR DE DATOS'])
ws.append([])
ws.append(['Comercial Andina del Pacífico S.A. (Ecuador)'])
ws.append([])
row('Tipo de Estado Financiero', *['Anual, Indiv.'] * 5)
row('Fecha final del período', *[datetime.date(y, 12, 31) for y in (2024, 2023, 2022, 2021, 2020)])
row('Unidades', *['Miles USD'] * 5)
ws.append([])
row('Estado de Resultados ', *Y)
row('    Total Ingreso Operativo', 52100.5, 49800.25, 47000, 44500, 41000)
row('    Ingresos netos por ventas', 51250.75, 48900.1, 46100, 43800, 40200)
row('    Costo de mercancías vendidas', -40000.2, -38000, -36500, -34000, -32000)
row('    Utilidad bruta', 11250.55, 10900.1, 9600, 9800, 8200)
row('    Gastos administrativos', -2000, -1900, -1800)
row('    Ganancia operativa (EBIT)', 5200.4, 5000, 4700, 4300, 3900)
row('    EBITDA', 6100.9, 5800, 5400, 5000, 4500)
row('    Resultado financiero', -820.35, -790, -700, -650, -600)
row('        Gastos financieros', -820.35, -790, -700, -650, -600)
row('            Gastos por intereses', -700.1, -690, -620, -600, -550)
row('    Ganancia (Pérdida) Neta', 3300.12, 3100, 2900, 2600, 2300)
ws.append([])
row('Balance General ', *Y)
row('    Activos Totales', 30000, 29000, 28000, 27000, 26000)
row('        Pasivos Totales', 20000, 19500, 19000, 18500, 18000)
row('            Créditos y préstamos no corrientes', 0, 1200.5, 2400, 3600, 4800)
row('            Deudas por Emisión de Bonos', 4500.25, 4000, 3500)
row('            Otros pasivos no corrientes', 300, 250, 200)
row('            Créditos y préstamos corrientes', 0, 900, 1500, 2100, 2700)
row('            Instrumentos financieros derivados', 150, 140, 130)
row('            Pasivos corrientes por impuesto a la renta', None, None, None, 80, 70)
row('        Efectivo o Equivalentes', 1800.5, 1500, 1400, 1300, 1200)
row('    Total de patrimonio y pasivos', 30000, 29000, 28000, 27000, 26000)
row('    Total de patrimonio', 10000, 9500, 9000, 8500, 8000)
row('        Patrimonio neto atribuible a los propietarios de la controladora', 10000, 9500, 9000, 8500, 8000)
row('Estado de Flujo de Efectivo ', *Y)
row('    Efectivo al inicio del período', 1500, 1400, 1300, 1200, 1100)
row('    Efectivo al final del período', 1800.5, 1500, 1400, 1300, 1200)
for c in ws['B']:
    if isinstance(c.value, datetime.date):
        c.number_format = 'yyyy-mm-dd'
for r in ws.iter_rows(min_row=6, max_row=6):
    for c in r[1:]:
        c.number_format = 'yyyy-mm-dd'
wb.save('emis-miles.xlsx')

# Verdad conocida: cifras en USD completos (miles ×1.000). Deuda = préstamos + bonos.
json.dump({
    'sales24': 51250750, 'sales23': 48900100, 'sales22': 46100000, 'sales21': 43800000,
    'ebitda': 6100900, 'finExp': 820350, 'cash': 1800500, 'netIncome': 3300120, 'equity': 10000000,
    'debt': 4500250
}, open('emis-miles.truth.json', 'w'), indent=1)
print('ok')
